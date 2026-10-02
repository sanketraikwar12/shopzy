import sys
import docx

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('doc/SHOPZY_Project_Synopsis (1).docx')
print("Document sections:", len(doc.sections))

for i, p in enumerate(doc.paragraphs):
    brs = p._p.xpath('.//w:br[@w:type="page"]')
    pb = p.paragraph_format.page_break_before
    if brs or pb:
        print(f"P{i:3d}: Page break (brs={len(brs)}, pb_before={pb}), text={repr(p.text[:50])}")

print("\nAll Paragraphs:")
for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        print(f"P{i:3d} [{p.alignment}]: {p.text}")
