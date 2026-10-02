import docx
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml, OxmlElement

doc = docx.Document()
section = doc.sections[0]
section.top_margin = Cm(2.54)
section.bottom_margin = Cm(2.54)
section.left_margin = Cm(2.54)
section.right_margin = Cm(2.54)
section.different_first_page_header_footer = True

# Helper: set cell shading
def set_cell_shading(cell, color_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}" w:val="clear"/>')
    cell._tc.get_or_add_tcPr().append(shd)

# Helper: set cell borders
def set_cell_borders(cell, top="none", bottom="none", left="none", right="none", 
                     color="CBD5E1", sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="{top}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="{left}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{bottom}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="{right}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)

# Helper: set cell margins (padding)
def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

# Helper: add callout box
def add_callout(doc, title, text, box_type="info"):
    colors = {
        "info": {"bg": "EFF6FF", "border": "2563EB", "title_color": RGBColor(0x1E, 0x40, 0xAF)},
        "tip": {"bg": "F0FDF4", "border": "16A34A", "title_color": RGBColor(0x15, 0x80, 0x3D)},
        "warning": {"bg": "FFFBEB", "border": "D97706", "title_color": RGBColor(0xB4, 0x53, 0x09)},
        "note": {"bg": "F8FAFC", "border": "64748B", "title_color": RGBColor(0x33, 0x41, 0x55)},
    }
    cfg = colors.get(box_type, colors["info"])
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    cell.width = Cm(16.0)
    set_cell_shading(cell, cfg["bg"])
    set_cell_borders(cell, left="single", color=cfg["border"], sz="24")
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run_t = p.add_run(f"{title}: ")
    run_t.bold = True
    run_t.font.name = "Calibri"
    run_t.font.size = Pt(10)
    run_t.font.color.rgb = cfg["title_color"]
    
    run_b = p.add_run(text)
    run_b.font.name = "Calibri"
    run_b.font.size = Pt(9.5)
    run_b.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    
    doc.add_paragraph("") # spacing

add_callout(doc, "ARCHITECTURAL CONSTRAINT", "MongoDB forbids combining $text and $geoNear in the same stage.", "warning")
doc.save("test_callout.docx")
print("Saved test callout successfully!")
