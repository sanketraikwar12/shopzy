import docx
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml, OxmlElement

doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.54)
section.bottom_margin = Cm(2.54)
section.left_margin = Cm(2.54)
section.right_margin = Cm(2.54)

# Test 1: Page number using fldSimple (which IS a valid child of w:p)
footer = section.footer
p_foot = footer.paragraphs[0]
p_foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = p_foot.add_run("Page ")
r.font.name = "Calibri"
r.font.size = Pt(9)

# In OOXML, w:fldSimple is a direct valid child of w:p
fldSimple_page = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
p_foot._p.append(fldSimple_page)

r2 = p_foot.add_run(" of ")
r2.font.name = "Calibri"
r2.font.size = Pt(9)

fldSimple_numpages = parse_xml(r'<w:fldSimple %s w:instr="NUMPAGES"/>' % nsdecls('w'))
p_foot._p.append(fldSimple_numpages)

# Add sample text
doc.add_heading("Test Document", level=1)
doc.add_paragraph("This is a test paragraph to verify Word opening.")

doc.save("test_valid.docx")
print("Saved test_valid.docx")
