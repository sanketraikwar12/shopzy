import docx
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml, OxmlElement

doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.54)
section.bottom_margin = Cm(2.54)
section.left_margin = Cm(2.54)
section.right_margin = Cm(2.54)
section.different_first_page_header_footer = True

# Helper: format cell strictly adhering to OOXML schema
def format_cell(cell, bg_color=None, border_color=None, border_sz="4", 
                left_border_sz=None, left_border_color=None,
                top_pad=100, bottom_pad=100, left_pad=120, right_pad=120,
                v_align="center"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    
    # 1. tcBorders
    if left_border_sz or border_color:
        borders_elm = OxmlElement('w:tcBorders')
        if border_color:
            for side in ['top', 'left', 'bottom', 'right']:
                b = OxmlElement(f'w:{side}')
                b.set(qn('w:val'), 'single')
                b.set(qn('w:sz'), str(border_sz))
                b.set(qn('w:space'), '0')
                b.set(qn('w:color'), border_color)
                borders_elm.append(b)
        elif left_border_sz:
            b_left = OxmlElement('w:left')
            b_left.set(qn('w:val'), 'single')
            b_left.set(qn('w:sz'), str(left_border_sz))
            b_left.set(qn('w:space'), '0')
            b_left.set(qn('w:color'), left_border_color or "2563EB")
            borders_elm.append(b_left)
            for side in ['top', 'bottom', 'right']:
                b = OxmlElement(f'w:{side}')
                b.set(qn('w:val'), 'none')
                borders_elm.append(b)
        tcPr.append(borders_elm)
        
    # 2. shd
    if bg_color:
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:fill'), bg_color)
        tcPr.append(shd)
        
    # 3. tcMar
    if top_pad or bottom_pad or left_pad or right_pad:
        tcMar = OxmlElement('w:tcMar')
        for side, val in [('top', top_pad), ('bottom', bottom_pad), ('left', left_pad), ('right', right_pad)]:
            m = OxmlElement(f'w:{side}')
            m.set(qn('w:w'), str(val))
            m.set(qn('w:type'), 'dxa')
            tcMar.append(m)
        tcPr.append(tcMar)
        
    # 4. vAlign
    if v_align:
        vAlign = OxmlElement('w:vAlign')
        vAlign.set(qn('w:val'), v_align)
        tcPr.append(vAlign)

def set_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    cantSplit = OxmlElement('w:cantSplit')
    trPr.append(cantSplit)

def set_row_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    trPr.append(tblHeader)
    cantSplit = OxmlElement('w:cantSplit')
    trPr.append(cantSplit)

# Header & Footer
header = section.header
p_head = header.paragraphs[0]
p_head.text = "SHOPZY — Software Requirements Specification"
p_head.alignment = WD_ALIGN_PARAGRAPH.RIGHT

footer = section.footer
p_foot = footer.paragraphs[0]
p_foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = p_foot.add_run("Confidential — Academic Major Project\tPage ")
fld1 = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
p_foot._p.append(fld1)
r2 = p_foot.add_run(" of ")
fld2 = parse_xml(r'<w:fldSimple %s w:instr="NUMPAGES"/>' % nsdecls('w'))
p_foot._p.append(fld2)

# Add a styled table
table = doc.add_table(rows=3, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = 'Table Grid'
set_row_header(table.rows[0])

headers = ["Col 1", "Col 2", "Col 3"]
for i, h in enumerate(headers):
    c = table.rows[0].cells[i]
    c.text = h
    format_cell(c, bg_color="1E3A8A", border_color="1E3A8A", v_align="center")
    c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    c.paragraphs[0].runs[0].font.bold = True

for r_idx in range(1, 3):
    set_row_cant_split(table.rows[r_idx])
    for c_idx in range(3):
        c = table.rows[r_idx].cells[c_idx]
        c.text = f"Data {r_idx},{c_idx}"
        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        format_cell(c, bg_color=bg, border_color="CBD5E1", v_align="center")

# Add a callout box
callout = doc.add_table(rows=1, cols=1)
callout.alignment = WD_TABLE_ALIGNMENT.CENTER
c_box = callout.rows[0].cells[0]
format_cell(c_box, bg_color="EFF6FF", left_border_sz="24", left_border_color="2563EB", top_pad=140, bottom_pad=140, left_pad=160, right_pad=160)
p = c_box.paragraphs[0]
r_title = p.add_run("NOTE: ")
r_title.bold = True
r_body = p.add_run("This is a callout box test.")

doc.save("test_full_elements.docx")
print("Saved test_full_elements.docx")
